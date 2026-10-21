import importlib.util
import io
import pytest


def parser():
    assert importlib.util.find_spec('app.services.document_text') is not None, 'Document extraction is missing'
    from app.services.document_text import extract_document, chunk_sections
    return extract_document, chunk_sections


def test_txt_locations_and_overlap():
    extract, chunk = parser()
    sections = extract((' '.join('word'+str(i) for i in range(1000))).encode(), 'txt')
    passages = chunk(sections)
    assert len(passages) == 3
    assert passages[0]['location'] == 'Paragraph 1'
    assert passages[0]['text'].split()[-50:] == passages[1]['text'].split()[:50]


@pytest.mark.parametrize('data,kind', [(b'', 'txt'), (b'\xff', 'txt'), (b'invalid','pdf'), (b'invalid','docx'), (b'abc','exe')])
def test_rejects_invalid_documents(data, kind):
    extract, _ = parser()
    with pytest.raises(ValueError): extract(data, kind)


def test_size_and_text_limits():
    extract, _ = parser()
    for text in (b'x'*(10*1024*1024+1), b'a'*200001):
        with pytest.raises(ValueError): extract(text, 'txt')


def test_docx_paragraph_locations():
    extract, _ = parser()
    from docx import Document
    doc=Document(); doc.add_paragraph('Evaporation makes water vapor.'); doc.add_paragraph('Rain returns water.')
    f=io.BytesIO(); doc.save(f)
    assert [s['location'] for s in extract(f.getvalue(),'docx')] == ['Paragraph 1','Paragraph 2']


def test_pdf_scanned_and_encrypted_rejection():
    extract, _ = parser()
    from pypdf import PdfWriter
    writer=PdfWriter(); writer.add_blank_page(width=100,height=100)
    f=io.BytesIO();writer.write(f)
    with pytest.raises(ValueError, match='searchable'): extract(f.getvalue(),'pdf')
    writer.encrypt('password'); f=io.BytesIO();writer.write(f)
    with pytest.raises(ValueError, match='encrypted'): extract(f.getvalue(),'pdf')


def test_pdf_preserves_page_numbers():
    extract,_=parser()
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
    writer=PdfWriter()
    for sentence in ('Water evaporates.', 'Vapor condenses.'):
        page=writer.add_blank_page(width=300,height=300)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
        stream=DecodedStreamObject();stream.set_data(f'BT /F1 12 Tf 10 200 Td ({sentence}) Tj ET'.encode())
        page[NameObject('/Contents')]=writer._add_object(stream)
    f=io.BytesIO();writer.write(f)
    assert [s['location'] for s in extract(f.getvalue(),'pdf')] == ['Page 1','Page 2']


def rag():
    assert importlib.util.find_spec('app.services.teacher_retrieval') is not None, 'Teacher retrieval is missing'
    from app.services.teacher_retrieval import build_context, GroundedResource, validate_grounded, InsufficientSources
    return build_context, GroundedResource, validate_grounded, InsufficientSources


def test_retrieval_bounds_and_explicit_pasted_source():
    build,_,_,_=rag()
    class Repo:
        def search(self,uid,ids,query):
            return [{'id':str(i),'document_id':'d','title':'Water','location':'Page 1','body':'Water '*1500} for i in range(8)]
    context=build({'source_document_ids':['d'],'topic':'Water','objectives':'Evaporation','source_material':'Water changes state.'},'u',Repo())
    assert sum(len(p['text']) for p in context) <= 16000
    assert context[0]['document_id'] is None
    assert context[0]['title']=='Lesson source material'
    assert len(context)<=8


def test_empty_source_search_is_explicit_failure():
    build,_,_,error=rag()
    class Repo:
        def search(self,*args): return []
    with pytest.raises(error): build_context_result=build({'source_document_ids':['d'],'topic':'Missing','objectives':'','source_material':''},'u',Repo())


def test_invalid_citations_and_insufficient_model_response():
    _,model,validate,error=rag()
    from app.domain.teacher import ResourceContent
    content=ResourceContent(title='Water',objectives=['Evaporation','Condensation','Rain'])
    for response in (model(sufficient=False,content=None,passage_ids=[]), model(sufficient=True,content=content,passage_ids=['invented']),model(sufficient=True,content=content,passage_ids=[])):
        with pytest.raises(error): validate(response,[{'passage_id':'real'}])


def test_no_documents_retains_existing_generation():
    build,_,_,_=rag()
    class Repo:
        def search(self,*args): raise AssertionError('No retrieval should run')
    assert build({'source_document_ids':[],'source_material':'','topic':'Water','objectives':''},'u',Repo())==[]


def test_document_routes_exist_and_upload_limit_precedes_parse():
    assert importlib.util.find_spec('app.api.v1.teacher_documents') is not None, 'Document API missing'
    from app.api.v1.teacher_documents import router
    paths={r.path for r in router.routes}
    assert '/teacher/documents/upload' in paths
    assert '/teacher/documents/{document_id}/retry' in paths
    assert '/teacher/lessons/{lesson_id}/sources' in paths


def test_worker_does_not_download_deleted_or_old_attempt():
    from app.services.document_worker import process_document
    class Repo:
        def document(self,*args): return {'status':'deleting','task_id':'new'}
        def download(self,*args): raise AssertionError('Deleted source must never be downloaded')
    assert process_document('d','u','old',Repo()) is False


def test_worker_publishes_only_current_attempt():
    from app.services.document_worker import process_document
    class Repo:
        published=None
        def document(self,*args): return {'status':'processing','task_id':'t','kind':'txt'}
        def is_user_deleting(self,*args): return False
        def download(self,*args): return b'Water evaporates.'
        def publish(self,d,u,t,chunks): self.published=chunks; return False
    repo=Repo()
    assert process_document('d','u','t',repo) is False
    assert repo.published[0]['location']=='Paragraph 1'


def test_library_keeps_interrupted_deletions_visible():
    from app.services.document_repository import DocumentRepository
    repo=DocumentRepository()
    calls=[]
    repo._request=lambda *args,**kwargs: calls.append(kwargs) or []
    repo.documents('owner')
    assert 'status' not in calls[0]['params'], 'Interrupted deletions must remain visible for retry'


@pytest.mark.parametrize('declared,parts', [(12*1024*1024,[]),(0,[b'a'*(6*1024*1024),b'b'*(6*1024*1024)])])
def test_oversized_upload_is_rejected_before_parser(declared,parts):
    import asyncio
    from app.api.upload_limit import DocumentUploadLimit
    sent=[];messages=iter(parts)
    async def parser_app(*args): raise AssertionError('Oversized upload reached parser')
    async def receive(): return {'type':'http.request','body':next(messages),'more_body':True}
    async def send(message): sent.append(message)
    asyncio.run(DocumentUploadLimit(parser_app)({'type':'http','path':'/api/v1/teacher/documents/upload','headers':[(b'content-length',str(declared).encode())]},receive,send))
    assert sent[0]['status']==413


def test_pdf_page_limit():
    extract,_=parser()
    from pypdf import PdfWriter
    writer=PdfWriter()
    for _ in range(201): writer.add_blank_page(width=100,height=100)
    f=io.BytesIO();writer.write(f)
    with pytest.raises(ValueError,match='200'): extract(f.getvalue(),'pdf')


def test_deletion_failure_can_be_retried_without_republishing():
    from app.services.document_repository import DocumentRepository
    from contextlib import nullcontext
    repo=DocumentRepository();events=[]
    repo.mark_deleting=lambda *args: events.append('deleting')
    repo.storage_lock=lambda *args:nullcontext()
    def unavailable(*args,**kwargs): raise RuntimeError('Storage unavailable')
    repo._storage=unavailable
    repo._request=lambda *args,**kwargs:events.append('removed')
    row={'id':'d','user_id':'u','object_path':'u/d.txt'}
    with pytest.raises(RuntimeError):repo.remove(row)
    assert events==['deleting']
    repo._storage=lambda *args,**kwargs:events.append('storage removed')
    repo.remove(row)
    assert events==['deleting','deleting','storage removed','removed']


def test_parser_does_not_log_document_fragments(caplog):
    extract,_=parser()
    with pytest.raises(ValueError):extract(b'private-document-fragment','pdf')
    assert not any('priva' in record.getMessage() for record in caplog.records)


def test_flashcard_prompt_requires_empty_question_and_answer_arrays(monkeypatch):
    import app.services.teacher_generation as service
    from app.services.teacher_retrieval import GroundedResource
    from app.domain.teacher import ResourceContent,Flashcard
    class Provider:
        def generate_structured(self,**kwargs):
            assert 'questions=[] and answer_key=[]' in kwargs['prompt']
            return GroundedResource(sufficient=True,content=ResourceContent(title='Water',flashcards=[Flashcard(front='Water',back='Liquid') for _ in range(12)]),passage_ids=['p'])
    monkeypatch.setattr(service,'get_text_provider',lambda:Provider())
    lesson={'subject':'science','grade':'5','topic':'Water','language':'en-US','objectives':'','source_material':''}
    service.TeacherGenerationService().generate(lesson,'flashcards',[{'passage_id':'p','document_id':'d','title':'Water','location':'Paragraph 1','text':'Water is liquid.'}])
