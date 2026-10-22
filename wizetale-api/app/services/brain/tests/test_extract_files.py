"""Offline tests for file extractors: photo, PDF, document, text, audio (fake BrainAI, no network)."""
import io
import json
import subprocess

import pytest
from PIL import Image, ImageDraw

from app.domain.brain.pipeline import ExtractInput, ExtractionError
from app.services.brain.ai.client import AIUnavailable
from app.services.brain.extract import audio, document, extract, images, pdf, photo, text
from app.services.brain.extract.chunking import LIMIT, group_paragraphs, split_long
from app.services.brain.extract.segments import Segment, Transcript, windows


class FakeCache:
    def __init__(self):
        self.data = {}

    @staticmethod
    def key(*parts):
        return json.dumps(parts, default=str)

    def get(self, op, key):
        return self.data.get((op, key))

    def put(self, op, key, value):
        self.data[(op, key)] = value


class FakeAI:
    """`generate(**kwargs)` returns the queued results in order; exceptions in the queue are raised."""

    def __init__(self, *results):
        self.cache, self.results, self.calls = FakeCache(), list(results), []

    def generate(self, **kwargs):
        self.calls.append(kwargs)
        result = self.results.pop(0) if self.results else AIUnavailable("down")
        if isinstance(result, BaseException):
            raise result
        return result


def expect_error(code, fn, *args):
    with pytest.raises(ExtractionError) as error:
        fn(*args)
    assert error.value.code == code
    return error.value


# ---------------------------------------------------------------- fixtures

def png(size=(800, 600)) -> bytes:
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    draw.text((40, 40), "Heteroskedasticity", fill="black")
    draw.rectangle((40, 100, 700, 110), fill="navy")
    out = io.BytesIO()
    image.save(out, format="PNG")
    return out.getvalue()


def _escape(line: str) -> str:
    return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages, *, size=(612, 792), title="Week 3: Heteroskedasticity") -> bytes:
    """A minimal, valid text PDF (Helvetica, one content stream per page) with a correct xref."""
    width, height = size
    n = len(pages)
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(f'{5 + 2 * i} 0 R' for i in range(n))}] /Count {n} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Title ({_escape(title)}) >>".encode(),
    ]
    for i, body in enumerate(pages):
        lines = [body[k:k + 80] for k in range(0, len(body), 80)] or [""]
        ops = " ".join(f"({_escape(line)}) Tj T*" for line in lines)
        stream = f"BT /F1 11 Tf 14 TL 40 {height - 60} Td {ops} ET".encode()
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {width} {height}] "
                       f"/Resources << /Font << /F1 3 0 R >> >> /Contents {6 + 2 * i} 0 R >>".encode())
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R /Info 4 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


PAGE_ONE = ("The OLS estimator is unbiased under the Gauss-Markov assumptions. "
            "Heteroskedasticity makes the usual standard errors wrong, so we use robust errors.")
PAGE_TWO = ("The White test regresses squared residuals on the regressors and their squares. "
            "A large LM statistic rejects homoskedasticity at the chosen level.")


# ---------------------------------------------------------------- helpers

def test_split_long_never_exceeds_limit_when_a_long_sentence_follows_text():
    text_in = ("word " * 230).strip() + ". " + "x" * 3000
    parts = split_long(text_in)
    assert all(len(p) <= LIMIT for p in parts)
    assert "".join(parts).replace(" ", "") == text_in.replace(" ", "")


def test_group_paragraphs_anchor_each_group_at_its_first_paragraph():
    blocks = ["Title", "a" * 1300, "short", "b" * 1300]
    chunks = group_paragraphs(blocks)
    assert [c.anchor.paragraph for c in chunks] == [1, 3]
    assert chunks[1].text.startswith("short")


def test_windows_do_not_inherit_a_bogus_end_from_an_earlier_segment():
    lines = [Segment(start_sec=0, end_sec=5000, text="bad timing.")]
    lines += [Segment(start_sec=60 + i * 4.0, end_sec=64 + i * 4.0, text=f"line {i}.") for i in range(30)]
    chunks = windows(lines)
    assert chunks[0].anchor.end_sec == 5000
    assert all(c.anchor.end_sec <= 200 for c in chunks[1:])
    assert all(c.anchor.end_sec >= c.anchor.start_sec for c in chunks)


def test_dispatch_rejects_unknown_kinds():
    expect_error("unsupported_type", extract, ExtractInput(kind="hologram"), FakeAI())


# ---------------------------------------------------------------- photo

def test_photo_ocr_returns_slide_anchor_and_thumbnail():
    ocr = photo.PhotoOCR(is_slide=True, slide_number=7, title="Heteroskedasticity", language="en-US",
                         text="Var(uᵢ | xᵢ) = σᵢ²\n\n  - robust standard errors  \n")
    ai = FakeAI(ocr)
    result = photo.extract(ExtractInput(kind="photo", data=png(), filename="slide.png"), ai)
    assert [(c.anchor.kind, c.anchor.slide) for c in result.chunks] == [("slide", 7)]
    assert "robust standard errors" in result.chunks[0].text
    assert result.title == "Heteroskedasticity" and result.language == "en-US"
    assert result.metadata == {"slide_detected": True, "ocr": "vision"}
    assert result.thumbnail[:3] == b"\xff\xd8\xff" and len(result.thumbnail) <= 200_000
    sent = ai.calls[0]
    assert sent["schema"] is photo.PhotoOCR and sent["parts"][0].mime_type == "image/jpeg"
    assert sent["variables"] == {"filename": "slide.png"}


def test_photo_illegible_is_unreadable():
    ai = FakeAI(photo.PhotoOCR(text="", legible=False))
    expect_error("unreadable", photo.extract, ExtractInput(kind="photo", data=png()), ai)


def test_photo_without_ai_or_tesseract_is_ai_unavailable(monkeypatch):
    monkeypatch.setattr(photo.shutil, "which", lambda name: None)
    expect_error("ai_unavailable", photo.extract, ExtractInput(kind="photo", data=png()), FakeAI(AIUnavailable("down")))


def test_photo_falls_back_to_tesseract(monkeypatch):
    monkeypatch.setattr(photo.shutil, "which", lambda name: "/usr/bin/tesseract")
    calls = []

    def run(args, **kwargs):
        calls.append(args)
        out = "List of available languages (2):\neng\nrus\n" if "--list-langs" in args else "Gauss-Markov\nBLUE\n"
        return subprocess.CompletedProcess(args, 0, stdout=out, stderr="")

    monkeypatch.setattr(photo.subprocess, "run", run)
    result = photo.extract(ExtractInput(kind="photo", data=png()), FakeAI(AIUnavailable("down")))
    assert result.metadata["ocr"] == "tesseract" and result.chunks[0].text == "Gauss-Markov BLUE"
    assert calls[1][-1] == "eng+rus"


def test_photo_tesseract_timeout_is_ai_unavailable(monkeypatch):
    monkeypatch.setattr(photo.shutil, "which", lambda name: "/usr/bin/tesseract")

    def run(args, **kwargs):
        raise subprocess.TimeoutExpired(args, kwargs.get("timeout"))

    monkeypatch.setattr(photo.subprocess, "run", run)
    expect_error("ai_unavailable", photo.extract, ExtractInput(kind="photo", data=png()), FakeAI(AIUnavailable("down")))


def test_photo_garbage_bytes_are_unreadable():
    ai = FakeAI()
    expect_error("unreadable", photo.extract, ExtractInput(kind="photo", data=b"\xff\xd8\xff\xe0 not a jpeg"), ai)
    expect_error("unreadable", photo.extract, ExtractInput(kind="photo", data=None), ai)
    assert ai.calls == []


@pytest.mark.filterwarnings("ignore::PIL.Image.DecompressionBombWarning")
def test_photo_with_too_many_pixels_is_too_large(monkeypatch):
    monkeypatch.setattr(images.Image, "MAX_IMAGE_PIXELS", 10_000)
    expect_error("too_large", photo.extract, ExtractInput(kind="photo", data=png((150, 100))), FakeAI())
    expect_error("too_large", photo.extract, ExtractInput(kind="photo", data=png((400, 300))), FakeAI())


# ---------------------------------------------------------------- pdf

def test_pdf_text_pages_get_page_anchors_title_and_thumbnail():
    ai = FakeAI()
    result = pdf.extract(ExtractInput(kind="pdf", data=make_pdf([PAGE_ONE, PAGE_TWO])), ai)
    assert result.page_count == 2
    assert [(c.anchor.kind, c.anchor.page) for c in result.chunks] == [("page", 1), ("page", 2)]
    assert "Gauss-Markov" in result.chunks[0].text and "White test" in result.chunks[1].text
    assert result.title == "Week 3: Heteroskedasticity"
    assert result.metadata == {"pdf_text": "text"}
    assert result.thumbnail[:3] == b"\xff\xd8\xff" and result.thumbnail_mime == "image/jpeg"
    assert ai.calls == []


@pytest.mark.parametrize("data", [b"%PDF-1.4\n garbage \x00\x01", b"", b"not a pdf at all"])
def test_corrupt_pdf_is_unreadable(data):
    expect_error("unreadable", pdf.extract, ExtractInput(kind="pdf", data=data), FakeAI())


def test_pdf_with_little_text_goes_through_ai_ocr_and_keeps_text_ocr_missed():
    ocr = pdf.PdfOCR(pages=[pdf.OCRPage(page=1, text="Scanned page about instrumental variables."),
                            pdf.OCRPage(page=2, text=""), pdf.OCRPage(page=9, text="ghost page")])
    ai = FakeAI(ocr)
    result = pdf.extract(ExtractInput(kind="pdf", data=make_pdf(["IV", "Slide 2 only"])), ai)
    assert result.metadata == {"pdf_text": "vision"}
    assert [(c.anchor.page, c.text) for c in result.chunks] == [
        (1, "Scanned page about instrumental variables."), (2, "Slide 2 only")]
    assert ai.calls[0]["parts"][0].mime_type == "application/pdf"


def test_scanned_pdf_without_ai_is_ai_unavailable():
    expect_error("ai_unavailable", pdf.extract, ExtractInput(kind="pdf", data=make_pdf(["", ""])),
                 FakeAI(AIUnavailable("down")))


def test_sparse_pdf_without_ai_keeps_its_text():
    result = pdf.extract(ExtractInput(kind="pdf", data=make_pdf(["Short note."])), FakeAI(AIUnavailable("down")))
    assert result.metadata == {"pdf_text": "text"} and result.chunks[0].text == "Short note."


def test_pdf_with_owner_password_only_is_still_read():
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(make_pdf([PAGE_ONE, PAGE_TWO]))))
    writer.encrypt(user_password="", owner_password="teacher", algorithm="RC4-128")
    out = io.BytesIO()
    writer.write(out)
    result = pdf.extract(ExtractInput(kind="pdf", data=out.getvalue()), FakeAI())
    assert [c.anchor.page for c in result.chunks] == [1, 2]

    locked = PdfWriter()
    locked.append(PdfReader(io.BytesIO(make_pdf([PAGE_ONE]))))
    locked.encrypt(user_password="secret", owner_password="teacher", algorithm="RC4-128")
    out = io.BytesIO()
    locked.write(out)
    expect_error("unreadable", pdf.extract, ExtractInput(kind="pdf", data=out.getvalue()), FakeAI())


def test_pdf_over_page_limit_is_too_large(monkeypatch):
    monkeypatch.setattr(pdf, "MAX_PAGES", 1)
    expect_error("too_large", pdf.extract, ExtractInput(kind="pdf", data=make_pdf([PAGE_ONE, PAGE_TWO])), FakeAI())


def test_pdf_thumbnail_of_a_very_tall_page_stays_small():
    thumbnail = pdf._thumbnail(make_pdf([PAGE_ONE], size=(200, 14400)))
    width, height = Image.open(io.BytesIO(thumbnail)).size
    assert width <= 480 and height <= 960


# ---------------------------------------------------------------- document

def docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Lecture 3: Variance")
    doc.add_paragraph("Variance is the expected squared deviation from the mean. " * 22)
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Estimator"
    table.cell(0, 2).text = "Formula"
    table.cell(1, 0).text = "Sample"
    table.cell(1, 1).text = "s²"
    table.cell(1, 2).text = "Σ(x-x̄)²/(n-1)"
    doc.add_paragraph("Bessel's correction removes the bias of the sample variance. " * 22)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_docx_keeps_reading_order_and_paragraph_anchors():
    result = document.extract(ExtractInput(kind="document", data=docx_bytes(), filename="lecture3.docx",
                                           mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
                              FakeAI())
    assert result.title == "Lecture 3: Variance"
    joined = " ".join(c.text for c in result.chunks)
    assert joined.index("Estimator | Formula") < joined.index("Bessel")
    assert "Estimator | Estimator" not in joined  # merged cell appears once
    table_chunk = next(c for c in result.chunks if "Estimator" in c.text)
    assert table_chunk.anchor.kind == "paragraph" and table_chunk.anchor.paragraph == 3
    assert result.chunks[0].anchor.paragraph == 1


def test_docx_that_expands_too_far_is_too_large(monkeypatch):
    monkeypatch.setattr(document, "MAX_EXPANDED", 1000)
    expect_error("too_large", document.extract, ExtractInput(kind="document", data=docx_bytes(), filename="x.docx"),
                 FakeAI())


def test_docx_member_that_lies_about_its_size_is_unreadable():
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"\x00" * 3_000_000)
    data = bytearray(out.getvalue())
    central = data.rindex(b"PK\x01\x02")
    data[central + 24:central + 28] = (1000).to_bytes(4, "little")  # declared uncompressed size
    expect_error("unreadable", document.extract, ExtractInput(kind="document", data=bytes(data), filename="x.docx"),
                 FakeAI())


def test_utf8_text_file_with_bom():
    data = "﻿Неделя 5\n\nДисперсия измеряет разброс.\n\nКовариация измеряет связь.".encode("utf-8")
    result = document.extract(ExtractInput(kind="document", data=data, filename="notes.txt"), FakeAI())
    assert result.title == "Неделя 5"
    assert result.chunks[0].anchor.kind == "paragraph" and result.chunks[0].text.startswith("Неделя 5")


@pytest.mark.parametrize("inp", [
    ExtractInput(kind="document", data="Привет".encode("cp1251"), filename="notes.txt"),
    ExtractInput(kind="document", data=b"not a zip", filename="notes.docx"),
    ExtractInput(kind="document", data=b"   \n\n  ", mime_type="text/plain"),
])
def test_unreadable_documents(inp):
    expect_error("unreadable", document.extract, inp, FakeAI())


# ---------------------------------------------------------------- text

def test_pasted_text_gets_title_and_paragraph_anchors():
    raw = "Week 7: Instrumental variables\n\n" + "\n\n".join(f"Paragraph {i}. " + "IV fixes endogeneity. " * 40
                                                              for i in range(3))
    result = text.extract(ExtractInput(kind="text", text=raw), FakeAI())
    assert result.title == "Week 7: Instrumental variables"
    assert all(c.anchor.kind == "paragraph" and c.anchor.paragraph >= 1 for c in result.chunks)
    assert result.chunks[-1].anchor.paragraph > 1


def test_text_from_bytes_and_empty_text():
    result = text.extract(ExtractInput(kind="text", data=b"One line about \xff covariance"), FakeAI())
    assert result.title is None and "covariance" in result.chunks[0].text
    expect_error("unreadable", text.extract, ExtractInput(kind="text", text=" \n\t "), FakeAI())


# ---------------------------------------------------------------- audio

def test_audio_transcript_becomes_timestamp_windows(monkeypatch):
    monkeypatch.setattr(audio, "_duration", lambda data: None)
    segments = [Segment(start_sec=i * 10.0, end_sec=i * 10.0 + 10, text=f"Sentence {i}.") for i in range(18)]
    ai = FakeAI(Transcript(title="Voice note", language="ru-RU", segments=segments))
    result = audio.extract(ExtractInput(kind="audio", data=b"OggS fake", mime_type="audio/ogg"), ai)
    assert all(c.anchor.kind == "timestamp" for c in result.chunks)
    assert result.chunks[0].anchor.start_sec == 0 and result.duration_sec == pytest.approx(180)
    assert result.language == "ru-RU" and result.metadata == {"transcript": "ai"}
    assert ai.calls[0]["parts"][0].mime_type == "audio/ogg"


def test_audio_errors(monkeypatch):
    monkeypatch.setattr(audio, "_duration", lambda data: None)
    expect_error("unreadable", audio.extract, ExtractInput(kind="audio", data=b"x"), FakeAI(Transcript()))
    expect_error("unreadable", audio.extract, ExtractInput(kind="audio", data=b""), FakeAI())
    expect_error("ai_unavailable", audio.extract, ExtractInput(kind="audio", data=b"x"), FakeAI(AIUnavailable("down")))
    monkeypatch.setattr(audio, "_duration", lambda data: 7300.0)
    ai = FakeAI()
    expect_error("too_large", audio.extract, ExtractInput(kind="audio", data=b"x"), ai)
    assert ai.calls == []
